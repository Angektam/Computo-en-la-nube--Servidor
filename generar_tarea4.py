from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

title = doc.add_heading('TAREA 4: Funcionalidades Serverless', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(13)

meta = doc.add_paragraph('Equipo: Los Rojos  |  Entrega: Lunes 26 de octubre de 2026')
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in meta.runs:
    run.font.size = Pt(9)

meta2 = doc.add_paragraph('Individual, media cuartilla (a mano o digital)')
meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in meta2.runs:
    run.font.size = Pt(9)
    run.italic = True

doc.add_paragraph()

p = doc.add_heading('Objetivo', level=2)
for run in p.runs:
    run.font.size = Pt(11)

p = doc.add_paragraph('Identificar 3 funcionalidades breves y esporádicas del proyecto inventario-cloud que se disparen por eventos.')
for run in p.runs:
    run.font.size = Pt(10)

doc.add_paragraph()

p = doc.add_heading('Para cada una indica:', level=2)
for run in p.runs:
    run.font.size = Pt(11)

items = ['Funcionalidad: qué hace', 'Evento: qué la dispara', 'Tiempo: cuánto tarda', 'Servicio AWS: Lambda, S3, etc.']
for item in items:
    p = doc.add_paragraph(item, style='List Bullet')
    for run in p.runs:
        run.font.size = Pt(10)

doc.add_paragraph()

p = doc.add_heading('Ejemplo (No uses, solo referencia):', level=2)
for run in p.runs:
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(128, 128, 128)

ejemplo_table = doc.add_table(rows=5, cols=2)
ejemplo_table.style = 'Light Grid Accent 1'

rows_data = [
    ('Funcionalidad', 'Reporte nocturno'),
    ('Evento', 'Horario 11 PM (EventBridge)'),
    ('Acción', 'Lambda comprime logs, guarda en S3'),
    ('Tiempo', '30 segundos'),
    ('Servicio', 'Lambda + S3 + EventBridge'),
]

for i, (label, value) in enumerate(rows_data):
    ejemplo_table.rows[i].cells[0].text = label
    ejemplo_table.rows[i].cells[1].text = value

doc.add_paragraph()

p = doc.add_heading('Tu respuesta (3 funcionalidades):', level=2)
for run in p.runs:
    run.font.size = Pt(11)
    run.bold = True

for num in range(1, 4):
    p = doc.add_paragraph(f'Funcionalidad {num}:')
    for run in p.runs:
        run.font.size = Pt(10)
        run.bold = True
    doc.add_paragraph('Evento: _________________________________________________________', style='List Bullet')
    doc.add_paragraph('Acción: _________________________________________________________', style='List Bullet')
    doc.add_paragraph('Tiempo: _________________________________________________________', style='List Bullet')
    doc.add_paragraph('Servicio AWS: ___________________________________________________', style='List Bullet')
    doc.add_paragraph()

p = doc.add_heading('Criterios', level=2)
criterios = ['✔ Funcionalidades reales del proyecto', '✔ Eventos específicos y realistas', '✔ Tiempos < 5 minutos', '✔ Procesos breves y esporádicos (NO continuos)', '✔ Media cuartilla máximo']
for criterio in criterios:
    doc.add_paragraph(criterio)

doc.add_paragraph()

p = doc.add_heading('Entrega', level=2)
for run in p.runs:
    run.font.bold = True
doc.add_paragraph('Por correo antes del lunes 26, 11:59 PM')
doc.add_paragraph('Asunto: "Tarea 4 - Serverless - [Tu Nombre]"')
doc.add_paragraph('Archivo: Tarea04_Serverless_[TuNombre].docx')
doc.add_paragraph('(Individual - cada miembro del equipo entrega su versión)')

output = 'C:\\Users\\angek\\OneDrive\\Documentos\\COMPUTO-EN-LA-NUVE\\Tarea04_Serverless.docx'
doc.save(output)
print(f'Tarea 4 creada: {output}')
